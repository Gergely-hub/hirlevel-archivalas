from __future__ import annotations

from pathlib import Path

from docx import Document

from app.db import Message, Source


def export_messages_to_docx(
    path: Path | str,
    *,
    source: Source | None,
    messages: list[Message],
    title: str = "",
) -> Path:
    path = Path(path)
    doc = Document()
    heading = title or (source.name if source else "Hírlevelek")
    doc.add_heading(heading, level=1)
    if source:
        doc.add_paragraph(f"Feladó: {source.email}")
    doc.add_paragraph(f"Darabszám: {len(messages)}")
    doc.add_paragraph("")

    for msg in messages:
        doc.add_heading(msg.received_at.replace("T", " "), level=2)
        doc.add_paragraph(msg.subject)
        doc.add_paragraph(msg.body_text or "")
        doc.add_paragraph("—" * 20)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
