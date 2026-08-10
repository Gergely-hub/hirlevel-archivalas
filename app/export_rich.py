"""Egy levél rich exportja: HTML (olvasó) → Word / PDF, képekkel."""

from __future__ import annotations

import html as html_lib
import os
import subprocess
import tempfile
from pathlib import Path
from typing import Optional

from bs4 import BeautifulSoup
from docx import Document

from app.db import Message, Source


def message_body_html(msg: Message) -> str:
    """Az olvasóban használt HTML; ha nincs, sima szöveg."""
    raw = (msg.body_html or "").strip()
    if raw:
        return raw
    text = html_lib.escape(msg.body_text or "(nincs tartalom)").replace("\n", "<br>\n")
    return f'<div style="font-family:Calibri,sans-serif;line-height:1.45">{text}</div>'


def build_export_document_html(msg: Message, *, source: Optional[Source] = None) -> str:
    """Önálló HTML dokumentum Word/PDF exporthoz."""
    body = message_body_html(msg)
    # Ha a törzs már teljes html, csak a body tartalmát vegyük
    lower = body.lower()
    if "<html" in lower:
        try:
            soup = BeautifulSoup(body, "lxml")
            body_el = soup.body
            body = body_el.decode_contents() if body_el else body
        except Exception:
            pass

    meta_bits = [html_lib.escape((msg.received_at or "").replace("T", " "))]
    if source:
        meta_bits.append(html_lib.escape(source.name or source.email or ""))
    meta = " · ".join(x for x in meta_bits if x)
    subject = html_lib.escape(msg.subject or "(nincs tárgy)")

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<title>{subject}</title>
<style>
body {{ font-family: Calibri, Candara, Segoe UI, sans-serif; line-height: 1.45;
       color: #1a1a1a; max-width: 780px; margin: 24px auto; padding: 0 16px; }}
h1 {{ font-size: 1.4em; margin: 0 0 8px; }}
.meta {{ color: #666; font-size: 0.95em; margin-bottom: 20px; }}
img {{ max-width: 100%; height: auto; }}
a {{ color: #0b57d0; }}
</style></head><body>
<h1>{subject}</h1>
<div class="meta">{meta}</div>
{body}
</body></html>
"""


def _body_fragment(html: str) -> str:
    """Teljes HTML dokumentumból a body tartalom (Word / beágyazás)."""
    if "<html" not in html.lower() and "<body" not in html.lower():
        return html
    try:
        soup = BeautifulSoup(html, "lxml")
        if soup.body:
            return soup.body.decode_contents()
    except Exception:
        pass
    return html


def export_message_to_docx(
    path: Path | str,
    msg: Message,
    *,
    source: Optional[Source] = None,
    html: Optional[str] = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    from htmldocx import HtmlToDocx

    doc = Document()
    doc.add_heading(msg.subject or "(nincs tárgy)", level=1)
    when = (msg.received_at or "").replace("T", " ")
    line = when
    if source:
        line = f"{when} — {source.name}" if when else source.name
    if line:
        doc.add_paragraph(line)
    doc.add_paragraph("")

    parser = HtmlToDocx()
    parser.options["images"] = True
    parser.options["tables"] = True
    fragment = _body_fragment((html or "").strip() or message_body_html(msg))
    parser.add_html_to_document(fragment, doc)
    doc.save(str(path))
    return path


def _find_chromium() -> Optional[Path]:
    env = os.environ
    names = [
        Path(env.get("PROGRAMFILES", r"C:\Program Files"))
        / "Microsoft"
        / "Edge"
        / "Application"
        / "msedge.exe",
        Path(env.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)"))
        / "Microsoft"
        / "Edge"
        / "Application"
        / "msedge.exe",
        Path(env.get("LOCALAPPDATA", ""))
        / "Microsoft"
        / "Edge"
        / "Application"
        / "msedge.exe",
        Path(env.get("PROGRAMFILES", r"C:\Program Files"))
        / "Google"
        / "Chrome"
        / "Application"
        / "chrome.exe",
        Path(env.get("LOCALAPPDATA", ""))
        / "Google"
        / "Chrome"
        / "Application"
        / "chrome.exe",
    ]
    for p in names:
        if p.is_file():
            return p
    return None


def export_message_to_pdf(
    path: Path | str,
    msg: Message,
    *,
    source: Optional[Source] = None,
    html: Optional[str] = None,
) -> Path:
    """HTML → PDF Edge/Chrome headless print-to-pdf-del (hű a hírlevél kinézetéhez)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    browser = _find_chromium()
    if browser is None:
        raise RuntimeError(
            "PDF-hez Edge vagy Chrome kell (headless nyomtatás). "
            "Telepítsd a Microsoft Edge-et, vagy ments Wordbe."
        )

    raw = (html or "").strip()
    if raw and "<html" in raw.lower():
        doc_html = raw
    elif raw:
        subject = html_lib.escape(msg.subject or "(nincs tárgy)")
        meta_bits = [(msg.received_at or "").replace("T", " ")]
        if source:
            meta_bits.append(source.name or source.email or "")
        meta = html_lib.escape(" · ".join(x for x in meta_bits if x))
        doc_html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{subject}</title>
<style>
body {{ font-family: Calibri, Candara, Segoe UI, sans-serif; line-height: 1.45;
       color: #1a1a1a; max-width: 780px; margin: 24px auto; padding: 0 16px; }}
img {{ max-width: 100%; height: auto; }}
</style></head><body>
<h1>{subject}</h1>
<div style="color:#666;margin-bottom:20px">{meta}</div>
{raw}
</body></html>"""
    else:
        doc_html = build_export_document_html(msg, source=source)

    out = path.resolve()
    with tempfile.TemporaryDirectory(prefix="hirlevel_pdf_") as td:
        html_path = Path(td) / "letter.html"
        html_path.write_text(doc_html, encoding="utf-8")
        uri = html_path.resolve().as_uri()
        cmd = [
            str(browser),
            "--headless=new",
            "--disable-gpu",
            "--no-first-run",
            "--no-pdf-header-footer",
            f"--print-to-pdf={out}",
            uri,
        ]
        proc = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        if not out.is_file() or out.stat().st_size < 50:
            err = (proc.stderr or proc.stdout or "").strip()[:400]
            raise RuntimeError(f"PDF készítés sikertelen. {err}")
    return out
